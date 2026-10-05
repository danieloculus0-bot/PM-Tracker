from flask import Flask, request, redirect, url_for, send_file
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
import sqlite3
from datetime import datetime, timedelta
from pathlib import Path
from io import BytesIO
import qrcode
import qrcode.image.svg
import json
import urllib.request
import urllib.parse

app = Flask(__name__)

BASE_DIR = Path.cwd()
EXCEL_FILE = BASE_DIR / "pm_data.xlsx"
DB_FILE = BASE_DIR / "pm_app.db"

SITE_CONFIG_FILE = BASE_DIR / "site_config.json"

DEFAULT_SITE_CONFIG = {
    "site_name": "PM Tracker",
    "weather_label": "Local Site",
    "weather_lat": 0.0,
    "weather_lon": 0.0,
    "weather_enabled": False,
    "host": "127.0.0.1",
    "port": 5000
}

def load_site_config():
    config = DEFAULT_SITE_CONFIG.copy()
    if SITE_CONFIG_FILE.exists():
        try:
            with open(SITE_CONFIG_FILE, "r", encoding="utf-8") as f:
                saved = json.load(f)
            if isinstance(saved, dict):
                config.update(saved)
        except Exception:
            pass
    return config

SITE_CONFIG = load_site_config()
SITE_NAME = str(SITE_CONFIG.get("site_name") or "PM Tracker").strip()

# Weather widget config
WEATHER_ENABLED = bool(SITE_CONFIG.get("weather_enabled", False))
WEATHER_LABEL = str(SITE_CONFIG.get("weather_label") or "Local Site")
WEATHER_LAT = float(SITE_CONFIG.get("weather_lat") or 0.0)
WEATHER_LON = float(SITE_CONFIG.get("weather_lon") or 0.0)
WEATHER_CACHE_MINUTES = 15

_weather_cache = {
    "fetched_at": None,
    "data": None
}


def db():
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = db()
    cur = conn.cursor()

    cur.execute("""
        CREATE TABLE IF NOT EXISTS completions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            machine_id TEXT NOT NULL,
            task_id TEXT NOT NULL,
            operator_name TEXT NOT NULL,
            completion_notes TEXT,
            completed_at TEXT NOT NULL
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS maintenance_requests (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            request_number TEXT,
            submitted_by TEXT NOT NULL,
            submitted_at TEXT NOT NULL,
            machine_id TEXT,
            machine_name TEXT,
            location TEXT,
            priority TEXT,
            description TEXT NOT NULL,
            status TEXT NOT NULL,
            assigned_to TEXT,
            resolution_notes TEXT,
            closed_at TEXT
        )
    """)

    conn.commit()
    conn.close()


def html_escape(value):
    text = "" if value is None else str(value)
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&#39;")
    )


def ensure_data_workbook():
    if EXCEL_FILE.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Machines"
    ws.append(["Machine ID", "Machine Name", "Department", "Location", "Manufacturer", "Model", "Serial Number", "Asset Tag", "Install Year", "Criticality", "Notes", "Active (Y/N)"])
    ws2 = wb.create_sheet("PM_Tasks")
    ws2.append(["Task ID", "Machine ID", "Task Name", "Task Description", "Frequency Unit (Days/Weeks/Months)", "Frequency Value", "Responsible Role", "Estimated Minutes", "Safety Notes", "Active (Y/N)"])
    ws3 = wb.create_sheet("Completion_Log")
    ws3.append(["Completion ID", "Machine ID", "Task ID", "Completed By", "Completion Date", "Completion Time", "Notes", "Pass/Fail"])
    wb.save(EXCEL_FILE)
    wb.close()


def get_workbook():
    ensure_data_workbook()
    return load_workbook(EXCEL_FILE, data_only=True)


def get_headers(ws):
    headers = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row=1, column=col).value
        if value is not None:
            headers[str(value).strip().lower()] = col
    return headers


def find_col(headers, candidates, required=True):
    for candidate in candidates:
        key = candidate.strip().lower()
        if key in headers:
            return headers[key]
    if required:
        raise KeyError(f"Missing required column. Expected one of: {', '.join(candidates)}")
    return None


def parse_dateish(value):
    if value in [None, ""]:
        return None
    if isinstance(value, datetime):
        return value
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(str(value), fmt)
        except Exception:
            pass
    try:
        return datetime.fromisoformat(str(value))
    except Exception:
        return None


def fmt_date(dt):
    return "" if not dt else dt.strftime("%Y-%m-%d")


def fmt_datetime(dt):
    return "" if not dt else dt.strftime("%Y-%m-%d %H:%M:%S")


def current_base_url():
    return f"http://{request.host}"


def get_workbook_data():
    return get_workbook()


def read_machines():
    wb = get_workbook_data()
    ws = wb["Machines"]
    headers = get_headers(ws)

    col_machine_id = find_col(headers, ["Machine ID", "MachineID"])
    col_machine_name = find_col(headers, ["Machine Name", "MachineName"])
    col_department = find_col(headers, ["Department"], required=False)
    col_location = find_col(headers, ["Location"], required=False)
    col_manufacturer = find_col(headers, ["Manufacturer"], required=False)
    col_model = find_col(headers, ["Model"], required=False)
    col_serial = find_col(headers, ["Serial Number", "SerialNumber"], required=False)
    col_criticality = find_col(headers, ["Criticality"], required=False)
    col_notes = find_col(headers, ["Notes"], required=False)

    machines = []
    for row in range(2, ws.max_row + 1):
        machine_id = str(ws.cell(row=row, column=col_machine_id).value or "").strip()
        machine_name = str(ws.cell(row=row, column=col_machine_name).value or "").strip()
        if not machine_id or not machine_name:
            continue

        machines.append({
            "machine_id": machine_id,
            "machine_name": machine_name,
            "department": str(ws.cell(row=row, column=col_department).value or "").strip() if col_department else "",
            "location": str(ws.cell(row=row, column=col_location).value or "").strip() if col_location else "",
            "manufacturer": str(ws.cell(row=row, column=col_manufacturer).value or "").strip() if col_manufacturer else "",
            "model": str(ws.cell(row=row, column=col_model).value or "").strip() if col_model else "",
            "serial_number": str(ws.cell(row=row, column=col_serial).value or "").strip() if col_serial else "",
            "criticality": str(ws.cell(row=row, column=col_criticality).value or "").strip() if col_criticality else "",
            "notes": str(ws.cell(row=row, column=col_notes).value or "").strip() if col_notes else "",
        })

    wb.close()
    return machines


def machine_lookup():
    machines = read_machines()
    return {m["machine_id"]: m for m in machines}


def read_tasks():
    wb = get_workbook_data()
    ws = wb["PM_Tasks"]
    headers = get_headers(ws)

    col_task_id = find_col(headers, ["Task ID", "TaskID"], required=False)
    col_machine_id = find_col(headers, ["Machine ID", "MachineID"])
    col_task_name = find_col(headers, ["Task Name", "TaskDescription"], required=False)
    col_task_desc = find_col(headers, ["Task Description", "TaskDescription"], required=False)
    col_freq_days = find_col(headers, ["FrequencyDays"], required=False)
    col_freq_value = find_col(headers, ["Frequency Value"], required=False)
    col_freq_unit = find_col(headers, ["Frequency Unit (Days/Weeks/Months)"], required=False)
    col_last_completed = find_col(headers, ["Last Completed", "LastCompleted"], required=False)
    col_next_due = find_col(headers, ["Next Due", "NextDue"], required=False)
    col_role = find_col(headers, ["Responsible Role"], required=False)
    col_est = find_col(headers, ["Estimated Minutes"], required=False)
    col_safety = find_col(headers, ["Safety Notes"], required=False)

    tasks = []
    task_counter = 1

    for row in range(2, ws.max_row + 1):
        machine_id = str(ws.cell(row=row, column=col_machine_id).value or "").strip()
        if not machine_id:
            continue

        task_id = str(ws.cell(row=row, column=col_task_id).value or "").strip() if col_task_id else ""
        if not task_id:
            task_id = f"TASK-{task_counter:03d}"

        task_name = str(ws.cell(row=row, column=col_task_name).value or "").strip() if col_task_name else ""
        task_description = str(ws.cell(row=row, column=col_task_desc).value or "").strip() if col_task_desc else ""

        if not task_name and task_description:
            task_name = task_description
        if not task_description and task_name:
            task_description = task_name

        frequency_days = None
        if col_freq_days:
            raw = ws.cell(row=row, column=col_freq_days).value
            try:
                frequency_days = int(float(raw))
            except Exception:
                frequency_days = None
        elif col_freq_value:
            raw_val = ws.cell(row=row, column=col_freq_value).value
            raw_unit = str(ws.cell(row=row, column=col_freq_unit).value or "").strip().lower() if col_freq_unit else "days"
            try:
                val = int(float(raw_val))
                if raw_unit in ["day", "days", ""]:
                    frequency_days = val
                elif raw_unit in ["week", "weeks"]:
                    frequency_days = val * 7
                elif raw_unit in ["month", "months"]:
                    frequency_days = val * 30
                elif raw_unit in ["quarter", "quarters"]:
                    frequency_days = val * 90
            except Exception:
                frequency_days = None

        last_completed = parse_dateish(ws.cell(row=row, column=col_last_completed).value) if col_last_completed else None
        next_due = parse_dateish(ws.cell(row=row, column=col_next_due).value) if col_next_due else None

        tasks.append({
            "task_id": task_id,
            "machine_id": machine_id,
            "task_name": task_name,
            "task_description": task_description,
            "frequency_days": frequency_days,
            "responsible_role": str(ws.cell(row=row, column=col_role).value or "").strip() if col_role else "",
            "estimated_minutes": str(ws.cell(row=row, column=col_est).value or "").strip() if col_est else "",
            "safety_notes": str(ws.cell(row=row, column=col_safety).value or "").strip() if col_safety else "",
            "last_completed_seed": last_completed,
            "next_due_seed": next_due,
        })
        task_counter += 1

    wb.close()
    return tasks


def latest_completion(machine_id, task_id):
    conn = db()
    row = conn.execute("""
        SELECT *
        FROM completions
        WHERE machine_id = ? AND task_id = ?
        ORDER BY datetime(completed_at) DESC
        LIMIT 1
    """, (machine_id, task_id)).fetchone()
    conn.close()
    return row


def recent_completions(limit=15):
    conn = db()
    rows = conn.execute("""
        SELECT *
        FROM completions
        ORDER BY datetime(completed_at) DESC
        LIMIT ?
    """, (limit,)).fetchall()
    conn.close()
    return rows


def add_completion(machine_id, task_id, operator_name, notes):
    conn = db()
    conn.execute("""
        INSERT INTO completions (machine_id, task_id, operator_name, completion_notes, completed_at)
        VALUES (?, ?, ?, ?, ?)
    """, (machine_id, task_id, operator_name, notes, datetime.now().isoformat(timespec="seconds")))
    conn.commit()
    conn.close()


def task_runtime(task):
    latest = latest_completion(task["machine_id"], task["task_id"])
    if latest:
        last_completed = parse_dateish(latest["completed_at"])
        completed_by = latest["operator_name"]
    else:
        last_completed = task["last_completed_seed"]
        completed_by = ""

    next_due = None
    if last_completed and task["frequency_days"]:
        next_due = last_completed + timedelta(days=task["frequency_days"])
    elif task["next_due_seed"]:
        next_due = task["next_due_seed"]

    today = datetime.now().date()
    status = "clear"
    label = "CLEAR"

    if next_due:
        delta = (next_due.date() - today).days
        if delta < 0:
            status = "overdue"
            label = "OVERDUE"
        elif delta == 0:
            status = "today"
            label = "DUE TODAY"
        elif delta <= 7:
            status = "soon"
            label = "DUE SOON"

    return {
        **task,
        "last_completed": last_completed,
        "next_due": next_due,
        "status": status,
        "status_label": label,
        "completed_by": completed_by,
    }


def machine_summary(machine, tasks):
    rt_tasks = [task_runtime(t) for t in tasks if t["machine_id"] == machine["machine_id"]]
    overdue = sum(1 for t in rt_tasks if t["status"] == "overdue")
    today = sum(1 for t in rt_tasks if t["status"] == "today")
    soon = sum(1 for t in rt_tasks if t["status"] == "soon")

    overall = "clear"
    overall_label = "CLEAR"
    if overdue > 0:
        overall = "overdue"
        overall_label = "OVERDUE"
    elif today > 0:
        overall = "today"
        overall_label = "DUE TODAY"
    elif soon > 0:
        overall = "soon"
        overall_label = "DUE SOON"

    due_dates = [t["next_due"] for t in rt_tasks if t["next_due"]]
    next_due = min(due_dates) if due_dates else None

    return {
        **machine,
        "tasks": rt_tasks,
        "overdue": overdue,
        "today": today,
        "soon": soon,
        "overall": overall,
        "overall_label": overall_label,
        "next_due": next_due,
    }


def all_machine_summaries():
    machines = read_machines()
    tasks = read_tasks()
    return [machine_summary(m, tasks) for m in machines]


def status_class(status):
    return {
        "overdue": "status-overdue",
        "today": "status-today",
        "soon": "status-soon",
        "clear": "status-clear",
    }.get(status, "status-clear")


def maintenance_status_class(status):
    status = (status or "").strip().lower()
    return {
        "open": "status-overdue",
        "in progress": "status-today",
        "closed": "status-clear",
    }.get(status, "status-soon")


def priority_class(priority):
    priority = (priority or "").strip().lower()
    return {
        "high": "status-overdue",
        "medium": "status-today",
        "low": "status-soon",
    }.get(priority, "status-soon")


def make_qr_svg(url):
    factory = qrcode.image.svg.SvgPathImage
    img = qrcode.make(url, image_factory=factory, box_size=8, border=2)
    stream = BytesIO()
    img.save(stream)
    svg = stream.getvalue().decode("utf-8")
    return svg


def weather_code_label(code):
    mapping = {
        0: "Clear sky",
        1: "Mainly clear",
        2: "Partly cloudy",
        3: "Overcast",
        45: "Fog",
        48: "Rime fog",
        51: "Light drizzle",
        53: "Drizzle",
        55: "Heavy drizzle",
        56: "Freezing drizzle",
        57: "Heavy freezing drizzle",
        61: "Light rain",
        63: "Rain",
        65: "Heavy rain",
        66: "Freezing rain",
        67: "Heavy freezing rain",
        71: "Light snow",
        73: "Snow",
        75: "Heavy snow",
        77: "Snow grains",
        80: "Rain showers",
        81: "Heavy rain showers",
        82: "Violent rain showers",
        85: "Snow showers",
        86: "Heavy snow showers",
        95: "Thunderstorm",
        96: "Thunderstorm hail",
        99: "Severe thunderstorm hail",
    }
    try:
        code = int(code)
    except Exception:
        return "Unknown"
    return mapping.get(code, f"Code {code}")


def weather_code_emoji(code, is_day=True):
    try:
        code = int(code)
    except Exception:
        return "🌤️"

    if code == 0:
        return "☀️" if is_day else "🌙"
    if code in [1, 2]:
        return "🌤️" if is_day else "☁️"
    if code == 3:
        return "☁️"
    if code in [45, 48]:
        return "🌫️"
    if code in [51, 53, 55, 56, 57]:
        return "🌦️"
    if code in [61, 63, 65, 66, 67, 80, 81, 82]:
        return "🌧️"
    if code in [71, 73, 75, 77, 85, 86]:
        return "❄️"
    if code in [95, 96, 99]:
        return "⛈️"
    return "🌤️"


def fetch_weather():
    if not WEATHER_ENABLED:
        return None

    now = datetime.now()
    fetched_at = _weather_cache.get("fetched_at")
    cached_data = _weather_cache.get("data")

    if fetched_at and cached_data:
        age = now - fetched_at
        if age.total_seconds() < WEATHER_CACHE_MINUTES * 60:
            return cached_data

    try:
        params = {
            "latitude": WEATHER_LAT,
            "longitude": WEATHER_LON,
            "current": ",".join([
                "temperature_2m",
                "apparent_temperature",
                "relative_humidity_2m",
                "weather_code",
                "is_day",
                "wind_speed_10m",
                "wind_gusts_10m",
                "precipitation",
            ]),
            "daily": ",".join([
                "weather_code",
                "temperature_2m_max",
                "temperature_2m_min",
                "precipitation_probability_max",
            ]),
            "temperature_unit": "fahrenheit",
            "wind_speed_unit": "mph",
            "precipitation_unit": "inch",
            "timezone": "auto",
            "forecast_days": 3,
        }
        url = "https://api.open-meteo.com/v1/forecast?" + urllib.parse.urlencode(params)
        with urllib.request.urlopen(url, timeout=5) as response:
            payload = json.loads(response.read().decode("utf-8"))

        current = payload.get("current", {})
        daily = payload.get("daily", {})

        forecast = []
        times = daily.get("time", [])
        codes = daily.get("weather_code", [])
        tmax = daily.get("temperature_2m_max", [])
        tmin = daily.get("temperature_2m_min", [])
        precip_prob = daily.get("precipitation_probability_max", [])

        for i in range(min(len(times), 3)):
            parsed = parse_dateish(times[i])
            forecast.append({
                "date": times[i],
                "label": parsed.strftime("%a") if parsed else times[i],
                "weather_code": codes[i] if i < len(codes) else None,
                "weather_text": weather_code_label(codes[i]) if i < len(codes) else "Unknown",
                "weather_icon": weather_code_emoji(codes[i], True) if i < len(codes) else "🌤️",
                "high": tmax[i] if i < len(tmax) else None,
                "low": tmin[i] if i < len(tmin) else None,
                "precip_probability": precip_prob[i] if i < len(precip_prob) else None,
            })

        result = {
            "location": WEATHER_LABEL,
            "current_temp": current.get("temperature_2m"),
            "feels_like": current.get("apparent_temperature"),
            "humidity": current.get("relative_humidity_2m"),
            "wind": current.get("wind_speed_10m"),
            "wind_gust": current.get("wind_gusts_10m"),
            "precipitation": current.get("precipitation"),
            "weather_code": current.get("weather_code"),
            "weather_text": weather_code_label(current.get("weather_code")),
            "weather_icon": weather_code_emoji(current.get("weather_code"), bool(current.get("is_day", 1))),
            "updated_at": current.get("time"),
            "forecast": forecast,
            "ok": True,
        }

        _weather_cache["fetched_at"] = now
        _weather_cache["data"] = result
        return result

    except Exception:
        result = {
            "location": WEATHER_LABEL,
            "ok": False,
            "forecast": [],
        }
        _weather_cache["fetched_at"] = now
        _weather_cache["data"] = result
        return result


def fmt_num(value, digits=0):
    if value in [None, ""]:
        return ""
    try:
        if digits == 0:
            return str(int(round(float(value))))
        return f"{float(value):.{digits}f}"
    except Exception:
        return str(value)


def weather_widget_html(weather):
    if not weather:
        return ""

    if not weather.get("ok"):
        return f"""
        <section class="panel weather-panel">
            <div class="panel-head">
                <div><h2>Local Weather</h2></div>
                <div class="muted">{html_escape(weather.get("location", ""))}</div>
            </div>
            <div class="empty-box">Weather unavailable right now.</div>
        </section>
        """

    forecast_cards = ""
    for day in weather.get("forecast", []):
        precip = ""
        if day.get("precip_probability") not in [None, ""]:
            precip = f'<div class="muted">Rain: {html_escape(fmt_num(day["precip_probability"]))}%</div>'

        forecast_cards += f"""
        <div class="forecast-card">
            <div class="forecast-day">{html_escape(day.get("label", ""))}</div>
            <div class="forecast-icon">{html_escape(day.get("weather_icon", "🌤️"))}</div>
            <div class="forecast-text">{html_escape(day.get("weather_text", ""))}</div>
            <div class="forecast-temps">
                <span>H {html_escape(fmt_num(day.get("high")))}°</span>
                <span>L {html_escape(fmt_num(day.get("low")))}°</span>
            </div>
            {precip}
        </div>
        """

    updated_line = ""
    updated_at = parse_dateish(weather.get("updated_at"))
    if updated_at:
        updated_line = updated_at.strftime("%Y-%m-%d %H:%M")

    return f"""
    <section class="panel weather-panel">
        <div class="panel-head">
            <div><h2>Local Weather</h2></div>
            <div class="muted">{html_escape(weather.get("location", ""))}</div>
        </div>

        <div class="weather-grid">
            <div class="weather-main">
                <div class="weather-now-icon">{html_escape(weather.get("weather_icon", "🌤️"))}</div>
                <div>
                    <div class="weather-now-temp">{html_escape(fmt_num(weather.get("current_temp")))}°F</div>
                    <div class="weather-now-text">{html_escape(weather.get("weather_text", ""))}</div>
                    <div class="muted">Feels like {html_escape(fmt_num(weather.get("feels_like")))}°F</div>
                </div>
            </div>

            <div class="weather-stats">
                <div class="mini-stat">
                    <div class="muted">Humidity</div>
                    <div class="mini-value">{html_escape(fmt_num(weather.get("humidity")))}%</div>
                </div>
                <div class="mini-stat">
                    <div class="muted">Wind</div>
                    <div class="mini-value">{html_escape(fmt_num(weather.get("wind")))} mph</div>
                </div>
                <div class="mini-stat">
                    <div class="muted">Gust</div>
                    <div class="mini-value">{html_escape(fmt_num(weather.get("wind_gust")))} mph</div>
                </div>
                <div class="mini-stat">
                    <div class="muted">Precip</div>
                    <div class="mini-value">{html_escape(fmt_num(weather.get("precipitation"), 2))}"</div>
                </div>
            </div>
        </div>

        <div class="forecast-grid">
            {forecast_cards}
        </div>

        <div class="muted" style="margin-top:16px;">Last weather update: {html_escape(updated_line) if updated_line else "Unknown"}</div>
    </section>
    """


def next_request_number():
    conn = db()
    row = conn.execute("SELECT id FROM maintenance_requests ORDER BY id DESC LIMIT 1").fetchone()
    conn.close()
    next_id = 1 if not row else int(row["id"]) + 1
    return f"MWO-{datetime.now().strftime('%Y%m%d')}-{next_id:03d}"


def add_maintenance_request(submitted_by, machine_id, location, priority, description):
    machines = machine_lookup()
    machine = machines.get(machine_id, {}) if machine_id else {}

    machine_name = machine.get("machine_name", "")
    machine_location = machine.get("location", "")
    final_location = (location or "").strip() or machine_location

    submitted_at = datetime.now().isoformat(timespec="seconds")
    request_number = next_request_number()

    conn = db()
    conn.execute("""
        INSERT INTO maintenance_requests (
            request_number,
            submitted_by,
            submitted_at,
            machine_id,
            machine_name,
            location,
            priority,
            description,
            status,
            assigned_to,
            resolution_notes,
            closed_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        request_number,
        submitted_by,
        submitted_at,
        machine_id,
        machine_name,
        final_location,
        priority,
        description,
        "Open",
        "",
        "",
        None
    ))
    conn.commit()
    row = conn.execute("SELECT last_insert_rowid() AS id").fetchone()
    conn.close()
    return int(row["id"])


def open_maintenance_requests(limit=50):
    conn = db()
    rows = conn.execute("""
        SELECT *
        FROM maintenance_requests
        WHERE lower(status) != 'closed'
        ORDER BY
            CASE lower(priority)
                WHEN 'high' THEN 1
                WHEN 'medium' THEN 2
                WHEN 'low' THEN 3
                ELSE 4
            END,
            datetime(submitted_at) DESC
        LIMIT ?
    """, (limit,)).fetchall()
    conn.close()
    return rows


def all_maintenance_requests(limit=200):
    conn = db()
    rows = conn.execute("""
        SELECT *
        FROM maintenance_requests
        ORDER BY datetime(submitted_at) DESC
        LIMIT ?
    """, (limit,)).fetchall()
    conn.close()
    return rows


def get_maintenance_request(req_id):
    conn = db()
    row = conn.execute("""
        SELECT *
        FROM maintenance_requests
        WHERE id = ?
    """, (req_id,)).fetchone()
    conn.close()
    return row


def count_open_maintenance_requests():
    conn = db()
    row = conn.execute("""
        SELECT COUNT(*) AS cnt
        FROM maintenance_requests
        WHERE lower(status) != 'closed'
    """).fetchone()
    conn.close()
    return int(row["cnt"])


def close_maintenance_request(req_id, assigned_to, resolution_notes):
    conn = db()
    conn.execute("""
        UPDATE maintenance_requests
        SET status = 'Closed',
            assigned_to = ?,
            resolution_notes = ?,
            closed_at = ?
        WHERE id = ?
    """, (
        assigned_to,
        resolution_notes,
        datetime.now().isoformat(timespec="seconds"),
        req_id
    ))
    conn.commit()
    conn.close()


def tabs_html(active="dashboard"):
    dash_class = "tab active" if active == "dashboard" else "tab"
    maint_class = "tab active" if active == "maintenance" else "tab"
    qr_class = "tab active" if active == "qr" else "tab"

    return f"""
    <div class="tabs">
        <a class="{dash_class}" href="/">Dashboard</a>
        <a class="{maint_class}" href="/maintenance">Maintenance Requests</a>
        <a class="{qr_class}" href="/qr">QR Codes</a>
        <a class="tab" href="/import">Import / Update Data</a>
        <a class="tab" href="/template.xlsx">Blank Template</a>
        <a class="tab green" href="/export.xlsx">Export Excel</a>
    </div>
    """


def layout(title, body):
    return f"""
    <!doctype html>
    <html>
    <head>
        <meta charset="utf-8">
        <meta name="viewport" content="width=device-width, initial-scale=1">
        <title>{html_escape(title)}</title>
        <style>
            :root {{
                --text: #f3f7ff;
                --muted: #a8b3cc;
            }}
            * {{ box-sizing: border-box; }}
            body {{
                margin: 0;
                font-family: Arial, Helvetica, sans-serif;
                color: var(--text);
                background:
                    radial-gradient(circle at top left, rgba(37,119,255,0.18), transparent 25%),
                    linear-gradient(180deg, #020713 0%, #031127 35%, #02102a 100%);
                min-height: 100vh;
            }}
            a {{ color: inherit; text-decoration: none; }}
            .shell {{
                width: min(1180px, calc(100vw - 28px));
                margin: 22px auto 40px auto;
                border: 1px solid rgba(114, 141, 189, 0.18);
                border-radius: 30px;
                background: linear-gradient(180deg, rgba(8,27,61,0.94), rgba(4,19,44,0.96));
                box-shadow: 0 20px 60px rgba(0,0,0,0.45);
                padding: 28px;
            }}
            .hero-card {{
                background: linear-gradient(180deg, rgba(21,43,84,0.80), rgba(17,38,73,0.45));
                border: 1px solid rgba(255,255,255,0.07);
                border-radius: 28px;
                padding: 26px;
                margin-bottom: 26px;
            }}
            .hero-title {{
                font-size: clamp(36px, 5vw, 62px);
                line-height: 0.95;
                font-weight: 900;
                margin: 0 0 14px 0;
            }}
            .hero-sub {{
                color: var(--muted);
                font-size: 24px;
                margin: 0;
            }}
            .tabs {{
                display: flex;
                gap: 16px;
                flex-wrap: wrap;
                margin-bottom: 28px;
            }}
            .tab, .btn {{
                border: none;
                border-radius: 18px;
                padding: 16px 24px;
                font-size: 18px;
                font-weight: 800;
                color: #fff;
                display: inline-block;
                cursor: pointer;
                box-shadow: inset 0 1px 0 rgba(255,255,255,0.08), 0 10px 24px rgba(0,0,0,0.30);
            }}
            .tab {{
                background: linear-gradient(180deg, #7f8492, #5d6170);
            }}
            .tab.active {{
                background: linear-gradient(180deg, #2f86ff, #1f5ee0);
            }}
            .tab.green {{
                background: linear-gradient(180deg, #20c667, #109a4f);
            }}
            .panel {{
                background: linear-gradient(180deg, rgba(8,24,56,0.94), rgba(5,19,45,0.96));
                border-radius: 28px;
                border: 1px solid rgba(255,255,255,0.08);
                padding: 24px;
                margin-bottom: 24px;
            }}
            .panel h2 {{
                margin: 0 0 8px 0;
                font-size: clamp(28px, 4vw, 48px);
                line-height: 1;
            }}
            .panel-head {{
                display: flex;
                justify-content: space-between;
                align-items: center;
                gap: 16px;
                flex-wrap: wrap;
                margin-bottom: 18px;
            }}
            .muted {{
                color: var(--muted);
            }}
            .kpi-grid {{
                display: grid;
                grid-template-columns: repeat(2, minmax(0, 1fr));
                gap: 20px;
                margin-bottom: 24px;
            }}
            .kpi {{
                background: linear-gradient(180deg, rgba(9,25,57,0.95), rgba(5,20,46,0.94));
                border-radius: 26px;
                padding: 24px;
                border: 1px solid rgba(255,255,255,0.08);
                min-height: 150px;
            }}
            .kpi .label {{
                color: var(--muted);
                font-size: 18px;
                text-transform: uppercase;
                letter-spacing: 2px;
                margin-bottom: 16px;
            }}
            .kpi .value {{
                font-size: clamp(42px, 5vw, 62px);
                font-weight: 900;
            }}
            .machine-grid {{
                display: grid;
                grid-template-columns: repeat(auto-fit, minmax(290px, 1fr));
                gap: 22px;
            }}
            .machine-card {{
                border-radius: 28px;
                padding: 24px;
                background: linear-gradient(180deg, rgba(7,23,53,0.95), rgba(5,18,42,0.97));
                border: 2px solid rgba(255,255,255,0.06);
            }}
            .machine-card.status-overdue {{ border-color: rgba(225,75,102,0.9); }}
            .machine-card.status-today {{ border-color: rgba(255,149,0,0.9); }}
            .machine-card.status-soon {{ border-color: rgba(211,163,44,0.9); }}
            .machine-card.status-clear {{ border-color: rgba(24,197,110,0.65); }}
            .machine-name {{
                font-size: 30px;
                font-weight: 900;
                margin-bottom: 6px;
            }}
            .machine-id {{
                color: var(--muted);
                font-size: 18px;
                margin-bottom: 20px;
            }}
            .pill {{
                display: inline-block;
                padding: 12px 20px;
                border-radius: 999px;
                font-weight: 900;
                font-size: 18px;
                margin-bottom: 20px;
                border: 1px solid rgba(255,255,255,0.12);
            }}
            .pill.status-overdue {{ background: rgba(225,75,102,0.17); color: #ffd5de; }}
            .pill.status-today {{ background: rgba(255,149,0,0.17); color: #ffe1b6; }}
            .pill.status-soon {{ background: rgba(211,163,44,0.17); color: #ffe9b0; }}
            .pill.status-clear {{ background: rgba(24,197,110,0.17); color: #c7ffd9; }}
            .stat-row {{
                display: grid;
                grid-template-columns: repeat(3, 1fr);
                gap: 16px;
                margin-bottom: 20px;
            }}
            .mini-stat {{
                background: rgba(255,255,255,0.03);
                border: 1px solid rgba(255,255,255,0.06);
                border-radius: 22px;
                padding: 18px;
                text-align: center;
            }}
            .mini-value {{
                font-size: 24px;
                font-weight: 900;
                margin-top: 6px;
            }}
            .action {{
                padding: 16px 24px;
                border-radius: 18px;
                font-size: 18px;
                font-weight: 900;
                display: inline-block;
                min-width: 130px;
                text-align: center;
                margin-right: 12px;
                margin-top: 10px;
                box-shadow: inset 0 1px 0 rgba(255,255,255,0.08), 0 10px 24px rgba(0,0,0,0.28);
            }}
            .action.blue {{ background: linear-gradient(180deg, #2f86ff, #1f5ee0); }}
            .action.silver {{ background: linear-gradient(180deg, #8b8f9e, #666b79); }}
            .action.green {{ background: linear-gradient(180deg, #20c667, #109a4f); }}
            table {{
                width: 100%;
                border-collapse: collapse;
            }}
            th, td {{
                padding: 16px 10px;
                border-bottom: 1px solid rgba(255,255,255,0.08);
                vertical-align: top;
                font-size: 18px;
            }}
            th {{
                color: var(--muted);
                text-transform: uppercase;
                letter-spacing: 2px;
                font-size: 16px;
                text-align: left;
            }}
            .empty-box {{
                border: 2px dashed rgba(255,255,255,0.12);
                border-radius: 24px;
                padding: 34px;
                color: var(--muted);
                text-align: center;
                font-size: 22px;
            }}
            .task-card {{
                border: 1px solid rgba(255,255,255,0.08);
                border-radius: 24px;
                padding: 20px;
                margin-bottom: 18px;
                background: rgba(255,255,255,0.02);
            }}
            .detail-line {{
                margin-bottom: 8px;
                color: var(--muted);
                font-size: 16px;
            }}
            input[type="text"], textarea, select {{
                width: 100%;
                padding: 16px 18px;
                border-radius: 18px;
                border: 1px solid rgba(255,255,255,0.12);
                background: rgba(255,255,255,0.04);
                color: #fff;
                font-size: 18px;
                outline: none;
            }}
            select option {{
                color: #000;
            }}
            textarea {{
                min-height: 110px;
                resize: vertical;
            }}
            .btn-row {{
                display: flex;
                gap: 14px;
                flex-wrap: wrap;
                margin-top: 10px;
            }}
            .qr-grid {{
                display: grid;
                grid-template-columns: repeat(auto-fit, minmax(320px, 1fr));
                gap: 24px;
            }}
            .qr-card {{
                border-radius: 28px;
                padding: 24px;
                background: linear-gradient(180deg, rgba(7,23,53,0.95), rgba(5,18,42,0.97));
                border: 1px solid rgba(255,255,255,0.08);
                text-align: center;
            }}
            .qr-wrap {{
                background: white;
                border-radius: 22px;
                padding: 18px;
                width: fit-content;
                margin: 14px auto 18px auto;
            }}
            .qr-wrap svg {{
                display: block;
                width: 240px;
                height: 240px;
            }}
            .qr-url {{
                color: var(--muted);
                font-size: 14px;
                word-break: break-all;
                margin-top: 10px;
            }}
            .weather-panel {{
                overflow: hidden;
            }}
            .weather-grid {{
                display: grid;
                grid-template-columns: 1.2fr 1fr;
                gap: 20px;
                align-items: stretch;
            }}
            .weather-main {{
                display: flex;
                align-items: center;
                gap: 20px;
                padding: 22px;
                border-radius: 24px;
                background: linear-gradient(180deg, rgba(13,35,76,0.92), rgba(7,24,52,0.95));
                border: 1px solid rgba(255,255,255,0.08);
            }}
            .weather-now-icon {{
                font-size: 64px;
                line-height: 1;
            }}
            .weather-now-temp {{
                font-size: clamp(42px, 5vw, 64px);
                font-weight: 900;
                line-height: 1;
                margin-bottom: 8px;
            }}
            .weather-now-text {{
                font-size: 24px;
                font-weight: 800;
                margin-bottom: 8px;
            }}
            .weather-stats {{
                display: grid;
                grid-template-columns: repeat(2, minmax(0, 1fr));
                gap: 14px;
            }}
            .forecast-grid {{
                display: grid;
                grid-template-columns: repeat(3, minmax(0, 1fr));
                gap: 16px;
                margin-top: 18px;
            }}
            .forecast-card {{
                border-radius: 22px;
                padding: 18px;
                background: rgba(255,255,255,0.03);
                border: 1px solid rgba(255,255,255,0.08);
                text-align: center;
            }}
            .forecast-day {{
                font-size: 20px;
                font-weight: 900;
                margin-bottom: 8px;
            }}
            .forecast-icon {{
                font-size: 34px;
                margin-bottom: 8px;
            }}
            .forecast-text {{
                font-size: 16px;
                font-weight: 700;
                min-height: 38px;
                margin-bottom: 10px;
            }}
            .forecast-temps {{
                display: flex;
                justify-content: center;
                gap: 14px;
                font-size: 18px;
                font-weight: 900;
                margin-bottom: 6px;
            }}
            .form-grid {{
                display: grid;
                grid-template-columns: repeat(2, minmax(0, 1fr));
                gap: 18px;
            }}
            .form-full {{
                grid-column: 1 / -1;
            }}
            .request-card {{
                border-radius: 24px;
                padding: 20px;
                background: rgba(255,255,255,0.03);
                border: 1px solid rgba(255,255,255,0.08);
                margin-bottom: 18px;
            }}
            .request-number {{
                font-size: 24px;
                font-weight: 900;
                margin-bottom: 8px;
            }}
            .request-desc {{
                font-size: 18px;
                margin: 10px 0 14px 0;
            }}
            @media print {{
                body {{
                    background: #fff;
                    color: #000;
                }}
                .shell {{
                    width: 100%;
                    margin: 0;
                    border: none;
                    box-shadow: none;
                    background: #fff;
                    color: #000;
                }}
                .tabs, .no-print {{
                    display: none !important;
                }}
                .qr-card {{
                    break-inside: avoid;
                    background: #fff;
                    color: #000;
                    border: 1px solid #ccc;
                }}
                .panel, .task-card, .request-card {{
                    background: #fff !important;
                    color: #000 !important;
                    border: 1px solid #ccc !important;
                    box-shadow: none !important;
                }}
                .qr-url, .muted, .detail-line {{
                    color: #333 !important;
                }}
                input, textarea, select, button {{
                    display: none !important;
                }}
            }}
            @media (max-width: 900px) {{
                .shell {{
                    padding: 18px;
                    border-radius: 24px;
                }}
                .tab, .btn {{
                    width: 100%;
                    text-align: center;
                }}
                .weather-grid, .forecast-grid, .form-grid {{
                    grid-template-columns: 1fr;
                }}
            }}
        </style>
    </head>
    <body>
        <div class="shell">{body}</div>
    </body>
    </html>
    """


@app.route("/")
def dashboard():
    init_db()
    summaries = all_machine_summaries()
    recent = recent_completions(15)
    open_requests = open_maintenance_requests(10)
    open_request_count = count_open_maintenance_requests()
    now = datetime.now()
    weather = fetch_weather()

    active_machines = len(summaries)
    overdue_tasks = sum(1 for s in summaries for t in s["tasks"] if t["status"] == "overdue")
    due_today = sum(1 for s in summaries for t in s["tasks"] if t["status"] == "today")
    due_soon = sum(1 for s in summaries for t in s["tasks"] if t["status"] == "soon")
    overdue_machines = sum(1 for s in summaries if s["overall"] == "overdue")
    completed_today = 0
    today = datetime.now().date()
    for r in recent_completions(500):
        dt = parse_dateish(r["completed_at"])
        if dt and dt.date() == today:
            completed_today += 1
    tasks_tracked = sum(len(s["tasks"]) for s in summaries)

    machine_cards = ""
    for s in summaries:
        machine_cards += f"""
        <div class="machine-card {status_class(s["overall"])}">
            <div class="machine-name">{html_escape(s["machine_name"])}</div>
            <div class="machine-id">{html_escape(s["machine_id"])}</div>
            <div class="pill {status_class(s["overall"])}">{html_escape(s["overall_label"])}</div>

            <div class="stat-row">
                <div class="mini-stat"><div class="mini-value">{s["overdue"]}</div><div class="muted">OVERDUE</div></div>
                <div class="mini-stat"><div class="mini-value">{s["today"]}</div><div class="muted">TODAY</div></div>
                <div class="mini-stat"><div class="mini-value">{s["soon"]}</div><div class="muted">SOON</div></div>
            </div>

            <div class="detail-line">Next Due: {html_escape(fmt_date(s["next_due"])) if s["next_due"] else "No tasks"}</div>

            <a class="action blue" href="/machine/{html_escape(s["machine_id"])}">Open</a>
        </div>
        """

    recent_html = ""
    if recent:
        task_map = {(t["machine_id"], t["task_id"]): t["task_description"] for s in summaries for t in s["tasks"]}
        machine_map = {s["machine_id"]: s["machine_name"] for s in summaries}
        recent_html += """
        <table>
            <tr>
                <th>Date</th>
                <th>Machine</th>
                <th>Task</th>
                <th>By</th>
            </tr>
        """
        for r in recent:
            recent_html += f"""
            <tr>
                <td>{html_escape(fmt_date(parse_dateish(r["completed_at"])))}</td>
                <td>{html_escape(machine_map.get(r["machine_id"], r["machine_id"]))}<br><span class="muted">{html_escape(r["machine_id"])}</span></td>
                <td>{html_escape(task_map.get((r["machine_id"], r["task_id"]), r["task_id"]))}</td>
                <td>{html_escape(r["operator_name"])}</td>
            </tr>
            """
        recent_html += "</table>"
    else:
        recent_html = '<div class="empty-box">No completed tasks yet.</div>'

    open_requests_html = ""
    if open_requests:
        open_requests_html += """
        <table>
            <tr>
                <th>Request</th>
                <th>Submitted</th>
                <th>Location</th>
                <th>Priority</th>
                <th>Status</th>
            </tr>
        """
        for r in open_requests:
            open_requests_html += f"""
            <tr>
                <td>
                    <a href="/maintenance/{r['id']}"><strong>{html_escape(r['request_number'] or f"REQ-{r['id']}")}</strong></a><br>
                    <span class="muted">{html_escape(r['machine_name'] or r['machine_id'] or 'General')}</span>
                </td>
                <td>{html_escape(r['submitted_by'])}<br><span class="muted">{html_escape(fmt_datetime(parse_dateish(r['submitted_at'])))}</span></td>
                <td>{html_escape(r['location'] or 'Not listed')}</td>
                <td><span class="pill {priority_class(r['priority'])}">{html_escape(r['priority'] or 'Medium')}</span></td>
                <td><span class="pill {maintenance_status_class(r['status'])}">{html_escape(r['status'])}</span></td>
            </tr>
            """
        open_requests_html += "</table>"
    else:
        open_requests_html = '<div class="empty-box">No open maintenance requests.</div>'

    body = f"""
    {tabs_html("dashboard")}

    <section class="hero-card">
        <div class="hero-title">{html_escape(SITE_NAME)}<br>PM Command<br>Center</div>
        <p class="hero-sub">Preventive maintenance, audit traceability, and real-time machine visibility.</p>
    </section>

    {weather_widget_html(weather)}

    <div class="kpi-grid">
        <div class="kpi"><div class="label">Current Date</div><div class="value">{now.strftime("%A,")}<br>{now.strftime("%B")}<br>{now.strftime("%d, %Y")}</div></div>
        <div class="kpi"><div class="label">Overdue Machines</div><div class="value">{overdue_machines}</div></div>
        <div class="kpi"><div class="label">Active Machines</div><div class="value">{active_machines}</div></div>
        <div class="kpi"><div class="label">Overdue Tasks</div><div class="value">{overdue_tasks}</div></div>
        <div class="kpi"><div class="label">Due Today</div><div class="value">{due_today}</div></div>
        <div class="kpi"><div class="label">Due Soon</div><div class="value">{due_soon}</div></div>
        <div class="kpi"><div class="label">Completed Today</div><div class="value">{completed_today}</div></div>
        <div class="kpi"><div class="label">Open Requests</div><div class="value">{open_request_count}</div></div>
    </div>

    <section class="panel">
        <div class="panel-head">
            <div><h2>Open Maintenance Requests</h2></div>
            <div>
                <a class="action green" href="/maintenance">Open Request Center</a>
            </div>
        </div>
        {open_requests_html}
    </section>

    <section class="panel">
        <div class="panel-head">
            <div><h2>Machine Status</h2></div>
            <div class="muted">QR printing and Excel export enabled.</div>
        </div>
        <div class="machine-grid">
            {machine_cards if machine_cards else '<div class="empty-box">No machines found.</div>'}
        </div>
    </section>

    <section class="panel">
        <div class="panel-head">
            <div><h2>Recent Completions</h2></div>
            <div class="muted">Most recent 15</div>
        </div>
        {recent_html}
    </section>
    """
    return layout(f"{SITE_NAME} PM Command Center", body)


@app.route("/machine/<machine_id>")
def machine(machine_id):
    init_db()
    summaries = all_machine_summaries()
    target = None
    for s in summaries:
        if s["machine_id"] == machine_id:
            target = s
            break

    if not target:
        return "<h1>Machine not found</h1>"

    task_blocks = ""
    for task in target["tasks"]:
        task_blocks += f"""
        <div class="task-card">
            <div style="font-size:24px;font-weight:900;margin-bottom:10px;">{html_escape(task["task_name"] or task["task_description"])}</div>
            <div class="pill {status_class(task["status"])}">{html_escape(task["status_label"])}</div>
            <div class="detail-line">Description: {html_escape(task["task_description"])}</div>
            <div class="detail-line">Frequency: {html_escape(str(task["frequency_days"]) + " days") if task["frequency_days"] else "Not set"}</div>
            <div class="detail-line">Last Completed: {html_escape(fmt_date(task["last_completed"])) if task["last_completed"] else "No history"}</div>
            <div class="detail-line">Next Due: {html_escape(fmt_date(task["next_due"])) if task["next_due"] else "Not scheduled"}</div>
            <div class="detail-line">Responsible Role: {html_escape(task["responsible_role"]) if task["responsible_role"] else "Not listed"}</div>
            <div class="detail-line">Estimated Minutes: {html_escape(task["estimated_minutes"]) if task["estimated_minutes"] else "Not listed"}</div>
            <div class="detail-line">Safety Notes: {html_escape(task["safety_notes"]) if task["safety_notes"] else "None"}</div>

            <form method="post" action="/complete-task">
                <input type="hidden" name="machine_id" value="{html_escape(target["machine_id"])}">
                <input type="hidden" name="task_id" value="{html_escape(task["task_id"])}">
                <input type="text" name="operator_name" placeholder="Operator name" required>
                <br><br>
                <textarea name="completion_notes" placeholder="Completion notes"></textarea>
                <div class="btn-row">
                    <button class="tab active" type="submit">Mark Complete</button>
                    <a class="action silver" href="/maintenance?machine_id={html_escape(target["machine_id"])}">Report Maintenance Issue</a>
                </div>
            </form>
        </div>
        """

    body = f"""
    {tabs_html("dashboard")}

    <section class="panel">
        <h2>{html_escape(target["machine_name"])}</h2>
        <div class="muted">{html_escape(target["machine_id"])}</div>
        <br>
        <div class="detail-line">Department: {html_escape(target["department"]) if target["department"] else "Not listed"}</div>
        <div class="detail-line">Location: {html_escape(target["location"]) if target["location"] else "Not listed"}</div>
        <div class="detail-line">Manufacturer: {html_escape(target["manufacturer"]) if target["manufacturer"] else "Not listed"}</div>
        <div class="detail-line">Model: {html_escape(target["model"]) if target["model"] else "Not listed"}</div>
        <div class="detail-line">Serial Number: {html_escape(target["serial_number"]) if target["serial_number"] else "Not listed"}</div>
        <div class="detail-line">Criticality: {html_escape(target["criticality"]) if target["criticality"] else "Not listed"}</div>
        <div class="detail-line">Notes: {html_escape(target["notes"]) if target["notes"] else "None"}</div>
    </section>

    <section class="panel">
        <div class="panel-head">
            <div><h2>PM Tasks</h2></div>
            <div>
                <a class="action green" href="/machine/{html_escape(target["machine_id"])}/clear-all">
                    Acknowledge / Clear All PMs
                </a>
            </div>
        </div>
        {task_blocks if task_blocks else '<div class="empty-box">No PM tasks found for this machine.</div>'}
    </section>
    """
    return layout(target["machine_name"], body)


@app.route("/machine/<machine_id>/clear-all", methods=["GET", "POST"])
def clear_all_machine_tasks(machine_id):
    init_db()
    summaries = all_machine_summaries()
    target = next((s for s in summaries if s["machine_id"] == machine_id), None)
    if not target:
        return "<h1>Machine not found</h1>", 404
    tasks = target["tasks"]
    if request.method == "POST":
        operator_name = request.form.get("operator_name", "").strip()
        completion_notes = request.form.get("completion_notes", "").strip()
        if not operator_name:
            return "Completed / acknowledged by is required.", 400
        note = "Bulk PM acknowledgement / clear all"
        if completion_notes:
            note += ": " + completion_notes
        for task in tasks:
            add_completion(target["machine_id"], task["task_id"], operator_name, note)
        return redirect(url_for("machine", machine_id=target["machine_id"]))
    task_rows = ""
    for task in tasks:
        task_rows += f"""
        <tr>
            <td>{html_escape(task["task_name"] or task["task_description"])}</td>
            <td><span class="pill {status_class(task["status"])}">{html_escape(task["status_label"])}</span></td>
            <td>{html_escape(fmt_date(task["last_completed"])) if task["last_completed"] else "No history"}</td>
        </tr>
        """
    body = f"""
    {tabs_html("dashboard")}
    <section class="panel">
        <h2>Acknowledge / Clear All PMs</h2>
        <div class="detail-line">Machine: <strong>{html_escape(target["machine_name"])}</strong></div>
        <div class="detail-line">Machine ID: {html_escape(target["machine_id"])}</div>
        <div class="detail-line">PM tasks to acknowledge: <strong>{len(tasks)}</strong></div>
        <br>
        <div class="empty-box">This records a completion acknowledgement for every PM task listed below. The person performing the acknowledgement is required.</div>
        <br>
        <form method="post">
            <label><strong>Completed / Acknowledged By</strong></label>
            <input type="text" name="operator_name" placeholder="Name" required>
            <br><br>
            <label><strong>Notes</strong></label>
            <textarea name="completion_notes" placeholder="Optional notes"></textarea>
            <br><br>
            <div class="btn-row">
                <button class="tab active" type="submit">Confirm and Clear All PMs</button>
                <a class="action silver" href="/machine/{html_escape(target["machine_id"])}">Cancel</a>
            </div>
        </form>
    </section>
    <section class="panel">
        <h2>Tasks Being Acknowledged</h2>
        <table>
            <tr><th>Task</th><th>Current Status</th><th>Last Completed</th></tr>
            {task_rows}
        </table>
    </section>
    """
    return layout("Acknowledge / Clear All PMs", body)


@app.route("/complete-task", methods=["POST"])
def complete_task():
    init_db()
    machine_id = request.form.get("machine_id", "").strip()
    task_id = request.form.get("task_id", "").strip()
    operator_name = request.form.get("operator_name", "").strip()
    completion_notes = request.form.get("completion_notes", "").strip()

    if not machine_id or not task_id or not operator_name:
        return "Missing required fields.", 400

    add_completion(machine_id, task_id, operator_name, completion_notes)
    return redirect(url_for("machine", machine_id=machine_id))


@app.route("/maintenance", methods=["GET", "POST"])
def maintenance():
    init_db()
    machines = read_machines()
    selected_machine_id = request.args.get("machine_id", "").strip()

    if request.method == "POST":
        submitted_by = request.form.get("submitted_by", "").strip()
        machine_id = request.form.get("machine_id", "").strip()
        location = request.form.get("location", "").strip()
        priority = request.form.get("priority", "Medium").strip()
        description = request.form.get("description", "").strip()

        if not submitted_by or not description:
            return "Submitted by and description are required.", 400

        req_id = add_maintenance_request(submitted_by, machine_id, location, priority, description)
        return redirect(url_for("maintenance_request_detail", req_id=req_id))

    options = '<option value="">General / Not Tied To Machine</option>'
    for m in machines:
        selected = "selected" if m["machine_id"] == selected_machine_id else ""
        label = f"{m['machine_name']} ({m['machine_id']})"
        options += f'<option value="{html_escape(m["machine_id"])}" {selected}>{html_escape(label)}</option>'

    request_rows = ""
    rows = all_maintenance_requests(200)
    if rows:
        request_rows += """
        <table>
            <tr>
                <th>Request</th>
                <th>Submitted</th>
                <th>Machine / Location</th>
                <th>Priority</th>
                <th>Status</th>
            </tr>
        """
        for r in rows:
            request_rows += f"""
            <tr>
                <td>
                    <a href="/maintenance/{r['id']}"><strong>{html_escape(r['request_number'] or f"REQ-{r['id']}")}</strong></a><br>
                    <span class="muted">{html_escape((r['description'] or '')[:80])}</span>
                </td>
                <td>{html_escape(r['submitted_by'])}<br><span class="muted">{html_escape(fmt_datetime(parse_dateish(r['submitted_at'])))}</span></td>
                <td>{html_escape(r['machine_name'] or r['machine_id'] or 'General')}<br><span class="muted">{html_escape(r['location'] or 'Not listed')}</span></td>
                <td><span class="pill {priority_class(r['priority'])}">{html_escape(r['priority'] or 'Medium')}</span></td>
                <td><span class="pill {maintenance_status_class(r['status'])}">{html_escape(r['status'])}</span></td>
            </tr>
            """
        request_rows += "</table>"
    else:
        request_rows = '<div class="empty-box">No maintenance requests submitted yet.</div>'

    body = f"""
    {tabs_html("maintenance")}

    <section class="panel">
        <div class="panel-head">
            <div><h2>Submit Maintenance Request</h2></div>
            <div class="muted">Create a work order for maintenance review.</div>
        </div>

        <form method="post" action="/maintenance">
            <div class="form-grid">
                <div>
                    <div class="muted" style="margin-bottom:8px;">Submitted By</div>
                    <input type="text" name="submitted_by" required>
                </div>
                <div>
                    <div class="muted" style="margin-bottom:8px;">Priority</div>
                    <select name="priority">
                        <option>Low</option>
                        <option selected>Medium</option>
                        <option>High</option>
                    </select>
                </div>
                <div>
                    <div class="muted" style="margin-bottom:8px;">Machine</div>
                    <select name="machine_id">
                        {options}
                    </select>
                </div>
                <div>
                    <div class="muted" style="margin-bottom:8px;">Location</div>
                    <input type="text" name="location" placeholder="Department, bay, cell, or area">
                </div>
                <div class="form-full">
                    <div class="muted" style="margin-bottom:8px;">Description</div>
                    <textarea name="description" placeholder="Describe the issue, symptoms, safety concern, downtime risk, or what needs repair." required></textarea>
                </div>
            </div>
            <div class="btn-row">
                <button class="tab active" type="submit">Submit Request</button>
            </div>
        </form>
    </section>

    <section class="panel">
        <div class="panel-head">
            <div><h2>Maintenance Request Log</h2></div>
            <div class="muted">Open and closed work orders.</div>
        </div>
        {request_rows}
    </section>
    """
    return layout("Maintenance Requests", body)


@app.route("/maintenance/<int:req_id>")
def maintenance_request_detail(req_id):
    init_db()
    r = get_maintenance_request(req_id)
    if not r:
        return "<h1>Maintenance request not found</h1>", 404

    close_form = ""
    if (r["status"] or "").lower() != "closed":
        close_form = f"""
        <section class="panel no-print">
            <h2>Close Work Order</h2>
            <form method="post" action="/maintenance/{r['id']}/close">
                <div class="form-grid">
                    <div>
                        <div class="muted" style="margin-bottom:8px;">Assigned / Closed By</div>
                        <input type="text" name="assigned_to" placeholder="Maintenance tech or responsible person">
                    </div>
                    <div></div>
                    <div class="form-full">
                        <div class="muted" style="margin-bottom:8px;">Resolution Notes</div>
                        <textarea name="resolution_notes" placeholder="What was repaired, adjusted, replaced, or verified?"></textarea>
                    </div>
                </div>
                <div class="btn-row">
                    <button class="tab active" type="submit">Close Request</button>
                    <a class="action silver" href="/maintenance/{r['id']}/print" target="_blank">Print Work Order</a>
                </div>
            </form>
        </section>
        """
    else:
        close_form = f"""
        <section class="panel no-print">
            <div class="btn-row">
                <a class="action silver" href="/maintenance/{r['id']}/print" target="_blank">Print Work Order</a>
            </div>
        </section>
        """

    body = f"""
    {tabs_html("maintenance")}

    <section class="panel">
        <div class="panel-head">
            <div>
                <h2>{html_escape(r['request_number'] or f"REQ-{r['id']}")}</h2>
                <div class="muted">Maintenance Request / Work Order</div>
            </div>
            <div>
                <span class="pill {maintenance_status_class(r['status'])}">{html_escape(r['status'])}</span>
                <span class="pill {priority_class(r['priority'])}">{html_escape(r['priority'] or 'Medium')}</span>
            </div>
        </div>

        <div class="detail-line">Submitted By: {html_escape(r['submitted_by'])}</div>
        <div class="detail-line">Submitted At: {html_escape(fmt_datetime(parse_dateish(r['submitted_at'])))}</div>
        <div class="detail-line">Machine ID: {html_escape(r['machine_id'] or 'Not listed')}</div>
        <div class="detail-line">Machine Name: {html_escape(r['machine_name'] or 'Not listed')}</div>
        <div class="detail-line">Location: {html_escape(r['location'] or 'Not listed')}</div>
        <div class="detail-line">Assigned To: {html_escape(r['assigned_to'] or 'Not assigned')}</div>
        <div class="detail-line">Closed At: {html_escape(fmt_datetime(parse_dateish(r['closed_at']))) if r['closed_at'] else 'Open'}</div>

        <br>
        <div style="font-size:22px;font-weight:900;margin-bottom:8px;">Issue Description</div>
        <div class="request-card">
            <div class="request-desc">{html_escape(r['description'])}</div>
        </div>

        <div style="font-size:22px;font-weight:900;margin-bottom:8px;">Resolution Notes</div>
        <div class="request-card">
            <div class="request-desc">{html_escape(r['resolution_notes'] or 'No resolution notes yet.')}</div>
        </div>

        <div class="btn-row no-print">
            <a class="action blue" href="/maintenance">Back To Requests</a>
            <a class="action silver" href="/maintenance/{r['id']}/print" target="_blank">Print Work Order</a>
        </div>
    </section>

    {close_form}
    """
    return layout("Maintenance Request", body)


@app.route("/maintenance/<int:req_id>/close", methods=["POST"])
def close_maintenance(req_id):
    init_db()
    r = get_maintenance_request(req_id)
    if not r:
        return "Request not found.", 404

    assigned_to = request.form.get("assigned_to", "").strip()
    resolution_notes = request.form.get("resolution_notes", "").strip()
    close_maintenance_request(req_id, assigned_to, resolution_notes)
    return redirect(url_for("maintenance_request_detail", req_id=req_id))


@app.route("/maintenance/<int:req_id>/print")
def print_maintenance(req_id):
    init_db()
    r = get_maintenance_request(req_id)
    if not r:
        return "Request not found.", 404

    body = f"""
    {tabs_html("maintenance")}

    <section class="panel">
        <div class="panel-head">
            <div>
                <h2>Work Order</h2>
                <div class="muted">{html_escape(r['request_number'] or f"REQ-{r['id']}")}</div>
            </div>
            <div class="no-print">
                <button class="tab active" onclick="window.print()">Print</button>
            </div>
        </div>

        <table>
            <tr><th style="width:240px;">Request Number</th><td>{html_escape(r['request_number'] or f"REQ-{r['id']}")}</td></tr>
            <tr><th>Submitted By</th><td>{html_escape(r['submitted_by'])}</td></tr>
            <tr><th>Submitted Date</th><td>{html_escape(fmt_datetime(parse_dateish(r['submitted_at'])))}</td></tr>
            <tr><th>Status</th><td>{html_escape(r['status'])}</td></tr>
            <tr><th>Priority</th><td>{html_escape(r['priority'] or 'Medium')}</td></tr>
            <tr><th>Machine ID</th><td>{html_escape(r['machine_id'] or '')}</td></tr>
            <tr><th>Machine Name</th><td>{html_escape(r['machine_name'] or '')}</td></tr>
            <tr><th>Location</th><td>{html_escape(r['location'] or '')}</td></tr>
            <tr><th>Assigned To</th><td>{html_escape(r['assigned_to'] or '')}</td></tr>
            <tr><th>Closed Date</th><td>{html_escape(fmt_datetime(parse_dateish(r['closed_at']))) if r['closed_at'] else ''}</td></tr>
            <tr><th>Description</th><td>{html_escape(r['description'])}</td></tr>
            <tr><th>Resolution Notes</th><td>{html_escape(r['resolution_notes'] or '')}</td></tr>
        </table>

        <br><br>
        <table>
            <tr>
                <th style="width:240px;">Maintenance Signature</th>
                <td style="height:60px;"></td>
            </tr>
            <tr>
                <th>Completion Date</th>
                <td></td>
            </tr>
        </table>
    </section>
    """
    return layout("Print Work Order", body)


@app.route("/qr")
def qr_page():
    init_db()
    summaries = all_machine_summaries()
    base_url = current_base_url()

    cards = ""
    for s in summaries:
        machine_url = f"{base_url}/machine/{s['machine_id']}"
        svg = make_qr_svg(machine_url)
        cards += f"""
        <div class="qr-card">
            <div class="machine-name" style="font-size:26px;">{html_escape(s["machine_name"])}</div>
            <div class="machine-id">{html_escape(s["machine_id"])}</div>
            <div class="qr-wrap">{svg}</div>
            <div class="qr-url">{html_escape(machine_url)}</div>
        </div>
        """

    body = f"""
    {tabs_html("qr")}

    <section class="panel">
        <div class="panel-head">
            <div>
                <h2>Machine QR Codes</h2>
                <div class="muted">Print these and place them on the machines.</div>
            </div>
        </div>
        <div class="btn-row no-print" style="margin-bottom:18px;">
            <button class="tab" onclick="window.print()">Print QR Codes</button>
        </div>
        <div class="qr-grid">
            {cards if cards else '<div class="empty-box">No machines found.</div>'}
        </div>
    </section>
    """
    return layout("Machine QR Codes", body)



@app.route("/template.xlsx")
def download_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Machines"
    ws.append(["Machine ID", "Machine Name", "Department", "Location", "Manufacturer", "Model", "Serial Number", "Asset Tag", "Install Year", "Criticality", "Notes", "Active (Y/N)"])
    ws2 = wb.create_sheet("PM_Tasks")
    ws2.append(["Task ID", "Machine ID", "Task Name", "Task Description", "Frequency Unit (Days/Weeks/Months)", "Frequency Value", "Responsible Role", "Estimated Minutes", "Safety Notes", "Active (Y/N)"])
    ws3 = wb.create_sheet("Completion_Log")
    ws3.append(["Completion ID", "Machine ID", "Task ID", "Completed By", "Completion Date", "Completion Time", "Notes", "Pass/Fail"])
    for wsx in (ws, ws2, ws3):
        for cell in wsx[1]:
            cell.font = Font(bold=True)
    output = BytesIO()
    wb.save(output)
    wb.close()
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="PM_Tracker_Import_Template.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _sheet_rows_as_dicts(ws):
    headers = [str(ws.cell(row=1, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(2, ws.max_row + 1):
        row = {}
        has_data = False
        for c, header in enumerate(headers, 1):
            if not header:
                continue
            value = ws.cell(row=r, column=c).value
            row[header] = value
            if value not in (None, ""):
                has_data = True
        if has_data:
            rows.append(row)
    return rows


def _ensure_sheet(wb, name, headers):
    if name in wb.sheetnames:
        ws = wb[name]
    else:
        ws = wb.create_sheet(name)
        ws.append(headers)
    current = [str(ws.cell(row=1, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
    for header in headers:
        if header not in current:
            ws.cell(row=1, column=ws.max_column + 1).value = header
            current.append(header)
    return ws


def _merge_sheet_by_key(dest_ws, source_rows, key_names):
    headers = [str(dest_ws.cell(row=1, column=c).value or "").strip() for c in range(1, dest_ws.max_column + 1)]
    header_map = {h: i + 1 for i, h in enumerate(headers) if h}
    existing = {}
    for r in range(2, dest_ws.max_row + 1):
        for key_name in key_names:
            if key_name not in header_map:
                continue
            value = str(dest_ws.cell(r, header_map[key_name]).value or "").strip()
            if value:
                existing[(key_name, value)] = r
    added = updated = 0
    for src in source_rows:
        target_row = None
        for key_name in key_names:
            value = str(src.get(key_name) or "").strip()
            if value and (key_name, value) in existing:
                target_row = existing[(key_name, value)]
                break
        if target_row is None:
            target_row = dest_ws.max_row + 1
            added += 1
        else:
            updated += 1
        for header, value in src.items():
            if header not in header_map:
                col = dest_ws.max_column + 1
                dest_ws.cell(row=1, column=col).value = header
                header_map[header] = col
            dest_ws.cell(row=target_row, column=header_map[header]).value = value
        for key_name in key_names:
            value = str(src.get(key_name) or "").strip()
            if value:
                existing[(key_name, value)] = target_row
    return added, updated


@app.route("/import", methods=["GET", "POST"])
def import_data():
    init_db()
    message = ""
    if request.method == "POST":
        upload = request.files.get("file")
        if not upload or not upload.filename:
            return "Select an .xlsx workbook.", 400
        if not upload.filename.lower().endswith(".xlsx"):
            return "Only .xlsx workbooks are supported.", 400
        try:
            source_wb = load_workbook(upload.stream, data_only=False)
            if EXCEL_FILE.exists():
                backup_name = BASE_DIR / ("pm_data_BACKUP_" + datetime.now().strftime("%Y%m%d_%H%M%S") + ".xlsx")
                backup_name.write_bytes(EXCEL_FILE.read_bytes())
                dest_wb = load_workbook(EXCEL_FILE)
            else:
                dest_wb = Workbook()
                dest_wb.remove(dest_wb.active)
            machine_headers = ["Machine ID", "Machine Name", "Department", "Location", "Manufacturer", "Model", "Serial Number", "Asset Tag", "Install Year", "Criticality", "Notes", "Active (Y/N)"]
            task_headers = ["Task ID", "Machine ID", "Task Name", "Task Description", "Frequency Unit (Days/Weeks/Months)", "Frequency Value", "Responsible Role", "Estimated Minutes", "Safety Notes", "Active (Y/N)"]
            log_headers = ["Completion ID", "Machine ID", "Task ID", "Completed By", "Completion Date", "Completion Time", "Notes", "Pass/Fail"]
            m_added = m_updated = t_added = t_updated = completions_added = 0
            if "Machines" in source_wb.sheetnames:
                rows = _sheet_rows_as_dicts(source_wb["Machines"])
                m_added, m_updated = _merge_sheet_by_key(_ensure_sheet(dest_wb, "Machines", machine_headers), rows, ["Machine ID"])
            if "PM_Tasks" in source_wb.sheetnames:
                rows = _sheet_rows_as_dicts(source_wb["PM_Tasks"])
                t_added, t_updated = _merge_sheet_by_key(_ensure_sheet(dest_wb, "PM_Tasks", task_headers), rows, ["Task ID"])
            if "Completion_Log" in source_wb.sheetnames:
                rows = _sheet_rows_as_dicts(source_wb["Completion_Log"])
                dest_ws = _ensure_sheet(dest_wb, "Completion_Log", log_headers)
                h = get_headers(dest_ws)
                mid_col = find_col(h, ["Machine ID"])
                tid_col = find_col(h, ["Task ID"])
                date_col = find_col(h, ["Completion Date"])
                by_col = find_col(h, ["Completed By"])
                existing_keys = set()
                for r in range(2, dest_ws.max_row + 1):
                    dt = parse_dateish(dest_ws.cell(r, date_col).value)
                    existing_keys.add((str(dest_ws.cell(r, mid_col).value or "").strip(), str(dest_ws.cell(r, tid_col).value or "").strip(), fmt_date(dt), str(dest_ws.cell(r, by_col).value or "").strip()))
                conn = db()
                for row in rows:
                    machine_id = str(row.get("Machine ID") or "").strip()
                    task_id = str(row.get("Task ID") or "").strip()
                    completed_by = str(row.get("Completed By") or "Imported").strip()
                    notes = str(row.get("Notes") or "").strip()
                    dt = parse_dateish(row.get("Completion Date"))
                    if not machine_id or not task_id or not dt:
                        continue
                    date_text = dt.strftime("%Y-%m-%d")
                    key = (machine_id, task_id, date_text, completed_by)
                    if key in existing_keys:
                        continue
                    dest_ws.append([row.get(header) for header in log_headers])
                    completed_at = dt.replace(hour=12, minute=0, second=0, microsecond=0).isoformat(timespec="seconds")
                    db_exists = conn.execute("SELECT 1 FROM completions WHERE machine_id=? AND task_id=? AND date(completed_at)=? AND operator_name=? LIMIT 1", (machine_id, task_id, date_text, completed_by)).fetchone()
                    if not db_exists:
                        conn.execute("INSERT INTO completions (machine_id, task_id, operator_name, completion_notes, completed_at) VALUES (?, ?, ?, ?, ?)", (machine_id, task_id, completed_by, notes, completed_at))
                    existing_keys.add(key)
                    completions_added += 1
                conn.commit()
                conn.close()
            _ensure_sheet(dest_wb, "Machines", machine_headers)
            _ensure_sheet(dest_wb, "PM_Tasks", task_headers)
            _ensure_sheet(dest_wb, "Completion_Log", log_headers)
            dest_wb.save(EXCEL_FILE)
            source_wb.close()
            dest_wb.close()
            message = f"Import complete. Machines added: {m_added}, updated: {m_updated}. Tasks added: {t_added}, updated: {t_updated}. Completion records added: {completions_added}."
        except Exception as exc:
            return f"Import failed: {html_escape(str(exc))}", 400
    body = f"""
    {tabs_html()}
    <section class="panel">
        <div class="panel-head"><div><h2>Import / Update PM Data</h2><div class="muted">Merge machines and tasks by ID without deleting existing records.</div></div></div>
        {"<div class='empty-box'>" + html_escape(message) + "</div>" if message else ""}
        <form method="post" enctype="multipart/form-data">
            <div style="margin:20px 0;"><input type="file" name="file" accept=".xlsx" required></div>
            <div class="btn-row">
                <button class="tab active" type="submit">Import / Update Data</button>
                <a class="action silver" href="/template.xlsx">Download Blank Template</a>
            </div>
        </form>
    </section>
    """
    return layout("Import / Update Data", body)


@app.route("/export.xlsx")
def export_excel():
    init_db()
    summaries = all_machine_summaries()
    recent = recent_completions(1000000)
    maint = all_maintenance_requests(500)

    wb = Workbook()
    dark_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    white_font = Font(color="FFFFFF", bold=True)

    ws1 = wb.active
    ws1.title = "Machine_Status"
    headers1 = ["Machine ID", "Machine Name", "Location", "Overall Status", "Overdue", "Due Today", "Due Soon", "Next Due"]
    ws1.append(headers1)
    for cell in ws1[1]:
        cell.fill = dark_fill
        cell.font = white_font

    for s in summaries:
        ws1.append([
            s["machine_id"],
            s["machine_name"],
            s["location"],
            s["overall_label"],
            s["overdue"],
            s["today"],
            s["soon"],
            fmt_date(s["next_due"])
        ])

    ws2 = wb.create_sheet("Overdue_Tasks")
    headers2 = ["Machine ID", "Machine Name", "Task ID", "Task Description", "Next Due"]
    ws2.append(headers2)
    for cell in ws2[1]:
        cell.fill = dark_fill
        cell.font = white_font

    for s in summaries:
        for t in s["tasks"]:
            if t["status"] == "overdue":
                ws2.append([
                    s["machine_id"],
                    s["machine_name"],
                    t["task_id"],
                    t["task_description"],
                    fmt_date(t["next_due"])
                ])

    ws3 = wb.create_sheet("Completion_History")
    headers3 = ["Completion Date", "Machine ID", "Task ID", "Operator", "Notes"]
    ws3.append(headers3)
    for cell in ws3[1]:
        cell.fill = dark_fill
        cell.font = white_font

    for r in recent:
        completed_dt = parse_dateish(r["completed_at"])
        ws3.append([
            completed_dt.strftime("%Y-%m-%d") if completed_dt else "",
            r["machine_id"],
            r["task_id"],
            r["operator_name"],
            r["completion_notes"]
        ])

    ws4 = wb.create_sheet("Maintenance_Requests")
    headers4 = ["Request Number", "Submitted By", "Submitted At", "Machine ID", "Machine Name", "Location", "Priority", "Status", "Assigned To", "Closed At", "Description", "Resolution Notes"]
    ws4.append(headers4)
    for cell in ws4[1]:
        cell.fill = dark_fill
        cell.font = white_font

    for r in maint:
        ws4.append([
            r["request_number"],
            r["submitted_by"],
            r["submitted_at"],
            r["machine_id"],
            r["machine_name"],
            r["location"],
            r["priority"],
            r["status"],
            r["assigned_to"],
            r["closed_at"],
            r["description"],
            r["resolution_notes"]
        ])

    for ws in [ws1, ws2, ws3, ws4]:
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val = "" if cell.value is None else str(cell.value)
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = min(max_len + 3, 45)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"pm_tracker_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        stream,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


if __name__ == "__main__":
    from waitress import serve
    init_db()
    host = str(SITE_CONFIG.get("host") or "127.0.0.1")
    port = int(SITE_CONFIG.get("port") or 5000)
    print(f"PM Tracker starting on http://{host}:{port}")
    serve(app, host=host, port=port, threads=8)