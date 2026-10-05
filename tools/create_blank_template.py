from pathlib import Path
from openpyxl import Workbook

out = Path(__file__).resolve().parents[1] / "PM_Tracker_Import_Template.xlsx"
wb = Workbook()
ws = wb.active
ws.title = "Machines"
ws.append(["Machine ID", "Machine Name", "Department", "Location", "Manufacturer", "Model", "Serial Number", "Asset Tag", "Install Year", "Criticality", "Notes", "Active (Y/N)"])
ws = wb.create_sheet("PM_Tasks")
ws.append(["Task ID", "Machine ID", "Task Name", "Task Description", "Frequency Unit (Days/Weeks/Months)", "Frequency Value", "Responsible Role", "Estimated Minutes", "Safety Notes", "Active (Y/N)"])
ws = wb.create_sheet("Completion_Log")
ws.append(["Completion ID", "Machine ID", "Task ID", "Completed By", "Completion Date", "Completion Time", "Notes", "Pass/Fail"])
wb.save(out)
wb.close()
print(out)
